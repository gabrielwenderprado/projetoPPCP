#!/usr/bin/env python3
"""Converte a aba planoAnual em demanda anual agrupada por modelo.

Fonte oficial: coluna C (Produto/modelo) e coluna D (MÊS).
"""
from __future__ import annotations

import json
import re
import sys
import unicodedata
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SOURCE = Path(sys.argv[1]) if len(sys.argv) > 1 else Path('/home/ubuntu/upload/CópiadeEXPLOSAO_02.10.xlsm')
OUTPUT = ROOT / 'data' / 'plano-anual.json'

MODEL_NAMES = ['18LDDI', '15LDDI', '13LDDI', '10S', '13AT', '10L', '13LDI']


def normalize(value: object) -> str:
    raw = unicodedata.normalize('NFKD', str(value or ''))
    raw = ''.join(char for char in raw if not unicodedata.combining(char))
    return re.sub(r'\s+', ' ', raw).strip().upper()


def model_for_product(product: object) -> str | None:
    name = normalize(product)
    if not name:
        return None
    # A coluna C possui descrições comerciais; estes testes mantêm o vínculo
    # explícito com os nomes dos modelos existentes na aba Pinos.
    if 'SKYCITY 10 S' in name:
        return '10S'
    if 'SKYCITY 10 L' in name:
        return '10L'
    if 'SKYCITY 10 HDOC' in name:
        return None
    if 'SKYCITY 13,5 AT' in name or 'SKYCITY 13.5 AT' in name or 'SKYCITY 13 AT' in name:
        return '13AT'
    if 'SKYCITY 13 LDI' in name:
        return '13LDI'
    if 'SKYCITY 15 LDDI' in name:
        return '15LDDI'
    if 'SKYCITY 18 LDDI' in name:
        return '18LDDI'
    return None


def month_value(value: object) -> tuple[int, int] | None:
    if isinstance(value, (datetime, date)):
        return value.year, value.month
    text = str(value or '').strip()
    match = re.search(r'(\d{1,2})[/-](\d{4})', text)
    if match:
        month, year = int(match.group(1)), int(match.group(2))
        return (year, month) if 1 <= month <= 12 else None
    return None


def main() -> int:
    if not SOURCE.exists():
        raise SystemExit(f'Planilha não encontrada: {SOURCE}')
    wb = load_workbook(SOURCE, read_only=True, data_only=True, keep_vba=True)
    sheet_name = next((name for name in wb.sheetnames if normalize(name) == 'PLANOANUAL'), None)
    if not sheet_name:
        raise SystemExit('Aba planoAnual não encontrada')
    ws = wb[sheet_name]
    models = {name: defaultdict(int) for name in MODEL_NAMES}
    products = Counter()
    rows = []
    for row in ws.iter_rows(min_row=2, max_col=4, values_only=True):
        pedido, cliente, product, month = (list(row) + [None] * 4)[:4]
        model = model_for_product(product)
        parsed = month_value(month)
        if not model or not parsed:
            continue
        year, month_number = parsed
        month_key = f'{month_number:02d}/{year}'
        models[model][month_key] += 1
        products[str(product).strip()] += 1
        rows.append({'pedido': str(pedido or '').strip(), 'cliente': str(cliente or '').strip(), 'produto': str(product or '').strip(), 'modelo': model, 'mes': month_key})
    payload = {
        'sourceFile': SOURCE.name,
        'sourceSheet': sheet_name,
        'generatedAt': date.today().isoformat(),
        'columns': {'model': 'C', 'month': 'D'},
        'models': {model: dict(sorted(months.items())) for model, months in models.items()},
        'products': dict(products),
        'rows': rows,
    }
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding='utf-8')
    print(f'PLANO_ANUAL_OK modelos={sum(bool(v) for v in models.values())} linhas={len(rows)} output={OUTPUT}')
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
