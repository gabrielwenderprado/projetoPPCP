"""Converte a aba única calfer em um snapshot consumido pelo dashboard."""
from pathlib import Path
from openpyxl import load_workbook
import json
import re
import sys

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / 'data' / 'calfer.json'


def text(value):
    return '' if value is None else str(value).strip()


def number(value):
    if value is None or value == '':
        return 0
    if isinstance(value, (int, float)):
        return float(value)
    raw = text(value).replace('.', '').replace(',', '.')
    try:
        return float(raw)
    except ValueError:
        return 0


def compact(value):
    value = float(value or 0)
    return int(value) if value.is_integer() else round(value, 6)


def model_name(value):
    value = text(value)
    return re.sub(r'^modelo\s+', '', value, flags=re.I).strip() or value


def read_sheet(ws, stock_label):
    models = []
    current = None
    for row_number, row in enumerate(ws.iter_rows(values_only=True), 1):
        values = list(row)
        first = text(values[0] if values else '')
        if first.lower().startswith('modelo'):
            current = {
                'name': model_name(first),
                'sourceLabel': first,
                'items': [],
            }
            models.append(current)
            continue
        if not current or first.lower() in {'código', 'codigo'} or not first:
            continue
        current['items'].append({
            'code': first,
            'description': text(values[1] if len(values) > 1 else ''),
            'quantityPerMachine': compact(number(values[2] if len(values) > 2 else 0)),
            stock_label: compact(number(values[3] if len(values) > 3 else 0)),
            'row': row_number,
        })
    models = [model for model in models if model['items']]
    items = []
    by_code = {}
    for model in models:
        for item in model['items']:
            code = item['code']
            if code not in by_code:
                by_code[code] = {
                    'code': code,
                    'description': item['description'],
                    'models': [],
                    'nextStock': 0,
                    'calferStock': 0,
                }
                items.append(by_code[code])
            summary = by_code[code]
            if model['name'] not in summary['models']:
                summary['models'].append(model['name'])
            summary[stock_label] = item.get(stock_label, 0)
    return models, items


def convert_single_calfer_sheet(ws, source_name):
    """Lê a aba única calfer: código, descrição, estoque Next, qtd/máquina e estoque Calfer."""
    models = []
    current = None
    for row_number, row in enumerate(ws.iter_rows(values_only=True), 1):
        values = list(row)
        first = text(values[0] if values else '')
        if first.lower().startswith('modelo'):
            current = {'name': model_name(first), 'sourceLabel': first, 'items': []}
            models.append(current)
            continue
        if not current or first.lower() in {'código', 'codigo'} or not first:
            continue
        current['items'].append({
            'code': first,
            'description': text(values[1] if len(values) > 1 else ''),
            'nextStock': compact(number(values[2] if len(values) > 2 else 0)),
            'quantityPerMachine': compact(number(values[3] if len(values) > 3 else 0)),
            'calferStock': compact(number(values[4] if len(values) > 4 else 0)),
            'row': row_number,
        })
    models = [model for model in models if model['items']]
    items = []
    by_code = {}
    for model in models:
        for item in model['items']:
            code = item['code']
            if code not in by_code:
                by_code[code] = {
                    'code': code,
                    'description': item['description'],
                    'models': [],
                    'nextStock': item['nextStock'],
                    'calferStock': item['calferStock'],
                }
                items.append(by_code[code])
            summary = by_code[code]
            if model['name'] not in summary['models']:
                summary['models'].append(model['name'])
            # A aba é a fonte da verdade; mantém a última ocorrência do código no modelo.
            summary['nextStock'] = item['nextStock']
            summary['calferStock'] = item['calferStock']
    return models, items

def convert(source):
    wb = load_workbook(source, read_only=True, data_only=True, keep_vba=True)
    names = {name.strip().lower(): name for name in wb.sheetnames}
    calfer_sheet = names.get('calfer')
    if not calfer_sheet:
        raise RuntimeError('A aba calfer é obrigatória.')
    calfer_models, items = convert_single_calfer_sheet(wb[calfer_sheet], Path(source).name)
    payload = {
        'sourceFile': Path(source).name,
        'generatedAt': __import__('datetime').date.today().isoformat(),
        'sourceSheets': {'calfer': calfer_sheet},
        'nextModels': [],
        'calferModels': calfer_models,
        'items': items,
        'meta': {
            'nextRows': sum(len(model['items']) for model in calfer_models),
            'calferRows': sum(len(model['items']) for model in calfer_models),
            'singleSheet': True,
        },
    }
    OUTPUT.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding='utf-8')
    print(f'CALFER_OK: {len(calfer_models)} modelos, {len(items)} códigos | {OUTPUT}')
    return payload
if __name__ == '__main__':
    if len(sys.argv) < 2:
        raise SystemExit('Uso: python scripts/convert_calfer.py /caminho/para/explosao.xlsm')
    convert(Path(sys.argv[1]).expanduser().resolve())
