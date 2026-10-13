#!/usr/bin/env python3
"""Converte planejamento_produtivo.xlsx em data/mapa-produtivo.json."""
from __future__ import annotations
from datetime import date, datetime
import json
import sys
from pathlib import Path
from openpyxl import load_workbook

STAGES = [('Corte e dobra', 5, 6), ('Solda bases', 7, 8), ('Subconjunto', 9, 10), ('Terceiro eixo', 11, 12), ('Encarroçamento', 13, 14), ('Máquina', 15, 16), ('Integração', 17, 18), ('Montagem final', 19, 20)]

def value(v):
    if isinstance(v, (datetime, date)): return v.isoformat()[:10]
    if v in ('#NUM!', '#VALUE!', '#N/A'): return None
    return v

def text(v): return '' if v is None else str(v).strip()

def norm_status(v): return text(v).upper().replace('Í', 'I').replace('É', 'E')

def extrair(source: Path, dest: Path) -> None:
    wb = load_workbook(source, data_only=True, read_only=True)
    ws = wb.active; rows = []; status_counts = {}; month_counts = {}; week_counts = {}
    for row in ws.iter_rows(min_row=4, values_only=True):
        if len(row) < 3 or not row[2]: continue
        status = text(row[25]); status_key = norm_status(status) or 'SEM STATUS'
        status_counts[status_key] = status_counts.get(status_key, 0) + 1
        month = value(row[24]); week = text(row[23])
        if month not in (None, ''): month_counts[str(month)] = month_counts.get(str(month), 0) + 1
        if week: week_counts[week] = week_counts.get(week, 0) + 1
        stages = []
        for name, planned_col, status_col in STAGES:
            stage_status = value(row[status_col])
            stages.append({'name': name, 'planned': value(row[planned_col]), 'statusValue': stage_status, 'completed': stage_status not in (None, '', 'n/a', 'N/A')})
        rows.append({'en': text(row[2]), 'pv': value(row[0]), 'pl': value(row[1]), 'client': value(row[3]), 'model': value(row[4]), 'family': value(row[28]), 'truckModel': value(row[30]), 'productionStatus': status or None, 'productionStatusNormalized': status_key, 'finalDate': value(row[22]), 'firstDeliveryDate': value(row[27]), 'week': week or None, 'month': month, 'vehicleArrival': value(row[29]), 'stages': stages, 'missingItems': [], 'missingItemsSource': 'calculado pelo sistema'})
    excluded = {'ENTREGUE', 'VITRINI'}; active = [r for r in rows if r['productionStatusNormalized'] not in excluded]
    payload = {'sourceFile': source.name, 'sourceSheet': ws.title, 'generatedAt': date.today().isoformat(), 'records': rows, 'summary': {'totalENs': len(rows), 'activeENs': len(active), 'excludedFromNiguri': len(rows)-len(active), 'byStatus': status_counts, 'byMonth': month_counts, 'byWeek': week_counts, 'excludedStatuses': sorted(excluded), 'stageRule': 'Status preenchido = etapa concluída'}, 'stageColumns': [{'name': n, 'plannedColumn': n, 'statusColumn': 'Status'} for n, _, _ in STAGES]}
    dest.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding='utf-8')
    print(f'MAPA_PRODUTIVO_OK {len(rows)} ENs | excluídas do NIGURI={len(rows)-len(active)} | {dest}')

if __name__ == '__main__':
    root = Path(__file__).resolve().parents[1]
    if len(sys.argv) < 2: raise SystemExit('Uso: convert_mapa_produtivo.py planejamento_produtivo.xlsx [saida.json]')
    extrair(Path(sys.argv[1]).expanduser().resolve(), Path(sys.argv[2]).resolve() if len(sys.argv)>2 else root/'data'/'mapa-produtivo.json')
