from pathlib import Path
from datetime import date
import json, sys
from openpyxl import load_workbook

def text(v): return '' if v is None else str(v).strip()
def num(v):
    try: return float(v or 0)
    except (TypeError, ValueError): return 0

def extrair(src, dest):
    wb=load_workbook(src, read_only=True, data_only=True, keep_vba=True); ws=wb['Programacao']
    selected=list(ws.iter_rows(min_col=27,max_col=32,values_only=True))
    names=[text(v) for v in selected[0]]
    models=[{'name':name,'column':27+i,'total':0,'days':[]} for i,name in enumerate(names) if name]
    for row_index, values in enumerate(selected[1:], start=2):
        for i,entry in enumerate(models):
            value=num(values[entry['column']-27]) if entry['column']-27 < len(values) else 0
            if value:
                entry['total'] += value
                entry['days'].append({'row':row_index,'quantity':int(value) if value.is_integer() else round(value,4)})
    for entry in models: entry['total']=int(entry['total']) if entry['total'].is_integer() else round(entry['total'],4)
    payload={'sourceFile':src.name,'sourceSheet':'Programacao','generatedAt':date.today().isoformat(),'columns':'AA:AF','models':models}
    dest.write_text(json.dumps(payload,ensure_ascii=False,indent=2),encoding='utf-8')
    print(f'PROGRAMACAO_OK {len(models)} modelos | {dest}')
if __name__=='__main__':
    root=Path(__file__).resolve().parents[1]; extrair(Path(sys.argv[1]), Path(sys.argv[2]) if len(sys.argv)>2 else root/'data/programacao-modelos.json')
