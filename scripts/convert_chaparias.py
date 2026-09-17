from pathlib import Path
from datetime import date
import json
import sys
from openpyxl import load_workbook

OUTPUT_DEFAULT = Path(__file__).resolve().parents[1] / 'data' / 'chaparias.json'

def text(value):
    return '' if value is None else str(value).strip().replace('\xa0', '')

def number(value):
    if isinstance(value, (int, float)):
        return float(value or 0)
    raw = text(value).replace('.', '').replace(',', '.')
    try:
        return float(raw) if raw else 0
    except ValueError:
        return 0

def compact(value):
    value = float(value or 0)
    return int(value) if value.is_integer() else round(value, 4)

def extrair(origem: Path, destino: Path) -> None:
    wb = load_workbook(origem, read_only=True, data_only=True, keep_vba=True)
    sheet_name = next((name for name in wb.sheetnames if name.strip().lower() == 'chaparias'), None)
    if not sheet_name:
        raise ValueError('A aba chaparias não foi encontrada na planilha.')
    ws = wb[sheet_name]
    headers = {text(ws.cell(2, c).value).lower(): c for c in range(1, ws.max_column + 1)}
    code_col = headers.get('código') or headers.get('codigo') or 1
    desc_col = headers.get('descrição') or headers.get('descricao') or 2
    stock_col = headers.get('estoque') or 4
    min_col = next((c for k, c in headers.items() if 'minimo' in k or 'mínimo' in k), 5)
    items = {}
    for row in range(3, ws.max_row + 1):
        code = text(ws.cell(row, code_col).value)
        if not code or code.lower() in {'código', 'codigo'}:
            continue
        desc = text(ws.cell(row, desc_col).value)
        stock = compact(number(ws.cell(row, stock_col).value))
        minimum = compact(number(ws.cell(row, min_col).value))
        item = items.setdefault(code, {'code': code, 'description': desc, 'unit': 'UN', 'stock': stock, 'minimum': minimum, 'modelNeeds': {'CHAPARIAS': 0}, 'lastMovement': 'não tem'})
        if not item['description'] and desc: item['description'] = desc
        if stock or not item['stock']: item['stock'] = stock
        if minimum or not item['minimum']: item['minimum'] = minimum
    output = list(items.values())
    for item in output:
        item['totalNeed'] = item['minimum']
        item['modelCount'] = 1
    output.sort(key=lambda item: (-(number(item['minimum']) - number(item['stock'])), item['code']))
    payload = {'sourceFile': origem.name, 'sourceSheet': sheet_name, 'generatedAt': date.today().isoformat(), 'models': ['CHAPARIAS'], 'items': output, 'meta': {'headerRow': 2, 'firstDataRow': 3, 'codeColumn': code_col, 'descriptionColumn': desc_col, 'stockColumn': stock_col, 'minimumColumn': min_col}}
    destino.parent.mkdir(parents=True, exist_ok=True)
    destino.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding='utf-8')
    print(f'CHAPARIAS_OK {len(output)} itens | {destino}')

if __name__ == '__main__':
    if len(sys.argv) not in {2, 3}: raise SystemExit('Uso: python convert_chaparias.py origem.xlsm [destino.json]')
    extrair(Path(sys.argv[1]), Path(sys.argv[2]) if len(sys.argv) == 3 else OUTPUT_DEFAULT)
